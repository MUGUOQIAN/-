from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.data_loader import load_json
from src.models import BillItem, CostBreakdown
from src.pricing.engine import price_item


@dataclass
class ColumnMap:
    """原清单列映射（只读，不覆盖原有价格列）"""

    category: int | None = None
    name: int | None = None
    description: int | None = None
    work_content: int | None = None
    unit: int | None = None
    quantity: int | None = None
    material_fee: int | None = None
    labor_fee: int | None = None
    measure_fee: int | None = None


@dataclass
class OutputColumnMap:
    """表格末尾追加的核算结果列（0-based index）"""

    material_fee: int
    labor_fee: int
    measure_fee: int


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\n", "")


def _find_column(headers: list[str], aliases: list[str], *, exact: bool = False) -> int | None:
    for idx, header in enumerate(headers):
        for alias in aliases:
            if exact:
                if header == alias:
                    return idx
            elif alias == header or alias in header or header in alias:
                return idx
    return None


def _read_headers(sheet: Worksheet, header_row: int) -> list[str]:
    return [
        _normalize_header(sheet.cell(header_row, col).value)
        for col in range(1, sheet.max_column + 1)
    ]


def build_column_map(sheet: Worksheet, config: dict) -> ColumnMap:
    header_row = config["header_row"]
    headers = _read_headers(sheet, header_row)
    columns = config["columns"]

    col_map = ColumnMap(
        category=_find_column(headers, columns["category"]),
        name=_find_column(headers, columns["name"]),
        description=_find_column(headers, columns["description"]),
        work_content=_find_column(headers, columns["work_content"]),
        unit=_find_column(headers, columns["unit"]),
        quantity=_find_column(headers, columns["quantity"]),
        material_fee=_find_column(headers, columns["material_fee"]),
        labor_fee=_find_column(headers, columns["labor_fee"]),
        measure_fee=_find_column(headers, columns["measure_fee"]),
    )

    if col_map.name is None:
        raise ValueError("未找到清单名称列，请检查表头是否包含：清单名称/项目名称")
    if col_map.quantity is None:
        raise ValueError("未找到工程量列，请检查表头是否包含：工程量/面积")
    return col_map


def ensure_output_columns(sheet: Worksheet, config: dict) -> OutputColumnMap:
    """在表格末尾追加核算结果列；若已存在则复用，避免重复追加。"""
    header_row = config["header_row"]
    headers = _read_headers(sheet, header_row)
    output_cfg = config["output_columns"]

    mat_name = output_cfg["material_fee"]
    labor_name = output_cfg["labor_fee"]
    measure_name = output_cfg["measure_fee"]

    mat_col = _find_column(headers, [mat_name], exact=True)
    labor_col = _find_column(headers, [labor_name], exact=True)
    measure_col = _find_column(headers, [measure_name], exact=True)

    if mat_col is not None and labor_col is not None and measure_col is not None:
        return OutputColumnMap(
            material_fee=mat_col,
            labor_fee=labor_col,
            measure_fee=measure_col,
        )

    start_col = sheet.max_column + 1
    sheet.cell(header_row, start_col, mat_name)
    sheet.cell(header_row, start_col + 1, labor_name)
    sheet.cell(header_row, start_col + 2, measure_name)

    return OutputColumnMap(
        material_fee=start_col - 1,
        labor_fee=start_col,
        measure_fee=start_col + 1,
    )


def _cell_value(sheet: Worksheet, row: int, col: int | None) -> str:
    if col is None:
        return ""
    value = sheet.cell(row, col + 1).value
    if value is None:
        return ""
    return str(value).strip()


def _parse_quantity(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"无法解析工程量: {value}") from exc


def read_bill_items(sheet: Worksheet, config: dict) -> list[BillItem]:
    col_map = build_column_map(sheet, config)
    items: list[BillItem] = []
    start_row = config["data_start_row"]

    for row in range(start_row, sheet.max_row + 1):
        name = _cell_value(sheet, row, col_map.name)
        if not name:
            continue

        quantity_raw = sheet.cell(row, col_map.quantity + 1).value if col_map.quantity is not None else 0
        items.append(
            BillItem(
                row=row,
                category=_cell_value(sheet, row, col_map.category),
                name=name,
                description=_cell_value(sheet, row, col_map.description),
                work_content=_cell_value(sheet, row, col_map.work_content),
                unit=_cell_value(sheet, row, col_map.unit),
                quantity=_parse_quantity(quantity_raw),
            )
        )
    return items


def write_costs(
    sheet: Worksheet,
    output_cols: OutputColumnMap,
    row: int,
    cost: CostBreakdown,
) -> None:
    """将核算结果写入表格末尾追加列，不覆盖原清单价格。"""
    if not cost.matched:
        sheet.cell(row, output_cols.material_fee + 1, "")
        sheet.cell(row, output_cols.labor_fee + 1, "")
        sheet.cell(row, output_cols.measure_fee + 1, "")
        return

    sheet.cell(row, output_cols.material_fee + 1, cost.material_total)
    sheet.cell(row, output_cols.labor_fee + 1, cost.labor_total)
    sheet.cell(row, output_cols.measure_fee + 1, cost.measure_fee)


def process_workbook(input_path: Path, output_path: Path, sheet_name: str | None = None) -> list[tuple[BillItem, CostBreakdown]]:
    config = load_json("data/config/excel-columns.json")
    wb = load_workbook(input_path)
    sheet = wb[sheet_name] if sheet_name else wb.active

    output_cols = ensure_output_columns(sheet, config)
    items = read_bill_items(sheet, config)
    results: list[tuple[BillItem, CostBreakdown]] = []

    for item in items:
        cost = price_item(item)
        write_costs(sheet, output_cols, item.row, cost)
        results.append((item, cost))

    wb.save(output_path)
    return results


def create_sample_workbook(path: Path) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "报价清单"

    headers = [
        "分部",
        "清单名称",
        "项目描述",
        "工作内容",
        "单位",
        "工程量",
        "材料费",
        "人工费",
        "措施费",
        "备注",
    ]
    sheet.append(headers)

    sheet.append(
        [
            "楼地面",
            "细石混凝土找平层-楼8 厚70mm",
            (
                "1.找平层厚度:详见设计图纸及招标文件\n"
                "2.混凝土强度等级:C20\n"
                "3.结合层:详见设计图纸及招标文件\n"
                "4.做法:1.钢筋混凝土板面除灰后刷纯水泥浆一道\n"
                "2.70厚C20细石混凝土垫层"
            ),
            "1.基层处理 2.找平层铺设 3.等其他全部相关工作内容",
            "m²",
            2399.49,
            75000,
            30000,
            12000,
            "原清单报价",
        ]
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
