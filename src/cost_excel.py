from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.data_loader import load_json


@dataclass
class CostItem:
    name: str
    spec: str = ""
    unit: str = ""
    price: float = 0.0
    item_type: str = "material"  # material | labor
    aliases: list[str] = field(default_factory=list)


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_column(headers: list[str], aliases: list[str]) -> int | None:
    for idx, header in enumerate(headers):
        for alias in aliases:
            if alias == header or alias in header or header in alias:
                return idx
    return None


def _parse_price(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("元", "")
    return float(text)


def _parse_type(value: str, config: dict) -> str:
    text = value.strip()
    for labor_kw in config["type_values"]["labor"]:
        if labor_kw in text:
            return "labor"
    return "material"


def load_cost_excel(path: Path, sheet_name: str | None = None) -> dict[str, CostItem]:
    """从成本 Excel 加载材料/人工单价，返回 {名称: CostItem}。"""
    config = load_json("data/config/cost-excel-columns.json")
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet: Worksheet = wb[sheet_name] if sheet_name else wb.active

    header_row = config["header_row"]
    headers = [
        _normalize(sheet.cell(header_row, col).value)
        for col in range(1, sheet.max_column + 1)
    ]
    cols = config["columns"]

    name_col = _find_column(headers, cols["name"])
    price_col = _find_column(headers, cols["price"])
    if name_col is None or price_col is None:
        wb.close()
        raise ValueError("成本 Excel 需包含「材料名称」和「单价」列")

    spec_col = _find_column(headers, cols["spec"])
    unit_col = _find_column(headers, cols["unit"])
    type_col = _find_column(headers, cols["type"])

    items: dict[str, CostItem] = {}
    for row in range(config["data_start_row"], sheet.max_row + 1):
        name = _normalize(sheet.cell(row, name_col + 1).value)
        if not name:
            continue
        price = _parse_price(sheet.cell(row, price_col + 1).value)
        if price <= 0:
            continue

        spec = _normalize(sheet.cell(row, spec_col + 1).value) if spec_col is not None else ""
        unit = _normalize(sheet.cell(row, unit_col + 1).value) if unit_col is not None else ""
        raw_type = _normalize(sheet.cell(row, type_col + 1).value) if type_col is not None else ""
        item_type = _parse_type(raw_type, config) if raw_type else (
            "labor" if "工日" in name or "人工" in name else "material"
        )

        entry = CostItem(name=name, spec=spec, unit=unit, price=price, item_type=item_type)
        items[name] = entry
        if spec:
            items[f"{name} {spec}"] = entry

    wb.close()
    return items


def create_sample_cost_workbook(path: Path) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "成本价"
    sheet.append(["材料名称", "规格", "单位", "单价", "类型"])
    sheet.append(["C20细石混凝土", "商品泵送", "m³", 415, "材料"])
    sheet.append(["C20预拌砂浆", "", "m³", 455, "材料"])
    sheet.append(["DS20预拌水泥砂浆", "", "m³", 485, "材料"])
    sheet.append(["素水泥浆", "", "m³", 520, "材料"])
    sheet.append(["稀水泥浆", "", "m³", 530, "材料"])
    sheet.append(["花岗岩石板", "20mm", "m²", 320, "材料"])
    sheet.append(["盲道花岗岩石砖", "300*300*25", "m²", 480, "材料"])
    sheet.append(["警示带花岗岩石材", "300*130", "m²", 350, "材料"])
    sheet.append(["水", "", "m³", 5, "材料"])
    sheet.append(["综合工日", "装饰工程", "工日", 135, "人工"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
